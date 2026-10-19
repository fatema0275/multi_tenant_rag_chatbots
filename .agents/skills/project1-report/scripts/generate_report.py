#!/usr/bin/env python3
"""
Project I (4IT31) report generator.

Turns a structured Markdown file (see resources/content_template.md) into a fully
formatted Word document that matches the department format:
  * Times New Roman everywhere (theme + styles + defaults are forced)
  * H1 18pt bold, H2 14pt bold, H3 12pt bold, body 12pt
  * justified, 1.5 line spacing, no page borders
  * header (left student IDs / right project title), footer (right page number)
  * Roman numerals for front matter, Arabic restarting at 1 for Chapter 1
  * auto Table of Contents / List of Figures / List of Tables via Word fields

The Cover Page and Certificate pages are NOT rebuilt. The original template
documents are copied in verbatim (via docxcompose) so their exact formatting,
fonts, spacing and logos are preserved.

Usage:
  python generate_report.py content.md -o Report.docx \
      --cover-template "report-content/Cover Page format-projectI.docx" \
      --certificate-template "report-content/Certificate format_projectI.docx"

Requires: python-docx, docxcompose, PyYAML.
"""
from __future__ import annotations

import argparse
import ctypes
import io
import os
import re
import sys
import zipfile
from ctypes import wintypes

try:
    import yaml
except ImportError:  # pragma: no cover
    yaml = None

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from docxcompose.composer import Composer
from copy import deepcopy

FONT = "Times New Roman"
H1_PT, H2_PT, H3_PT, BODY_PT = 18, 14, 12, 12
LINE = 1.5
PLACEHOLDER = "Right-click here and choose 'Update Field' to generate."

DEFAULT_COVER = os.path.join("report-content", "Cover Page format-projectI.docx")
DEFAULT_CERT = os.path.join("report-content", "Certificate format_projectI.docx")


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #
def _set_font(run, size=None, bold=None, italic=None, name=FONT, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def _spacing(par, line=LINE, before=None, after=None):
    pf = par.paragraph_format
    pf.line_spacing = line
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)


def _add_field(par, instr, placeholder=PLACEHOLDER):
    r1 = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._element.append(fc)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._element.append(it)
    r3 = par.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._element.append(sep)
    r4 = par.add_run(placeholder)
    _set_font(r4, size=BODY_PT)
    r5 = par.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._element.append(end)


def _add_tc_field(par, text, flag):
    run = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    run._element.append(fc)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f' TC "{text}" \\f {flag} \\l 1 '
    run._element.append(it)
    run2 = par.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run2._element.append(end)


def _page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def _pgnum(section, fmt=None, start=None):
    sectpr = section._sectPr
    old = sectpr.find(qn("w:pgNumType"))
    if old is not None:
        sectpr.remove(old)
    el = OxmlElement("w:pgNumType")
    if fmt:
        el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))
    sectpr.append(el)


def _content_width_in(section):
    return (section.page_width - section.left_margin - section.right_margin) / 914400


def _no_borders(table):
    tblpr = table._tbl.tblPr
    for el in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(el)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto")
        borders.append(e)
    tblpr.append(borders)


def _build_header(section, ids, title, width_in):
    """Two column borderless table: IDs far left, project title far right."""
    header = section.header
    header.is_linked_to_previous = False
    el = header._element
    for child in list(el):
        if child.tag in (qn("w:p"), qn("w:tbl")):
            el.remove(child)

    c0w = 1.55
    c1w = max(2.0, width_in - c0w)
    tbl = header.add_table(rows=1, cols=2, width=Inches(width_in))
    tbl.autofit = False
    _no_borders(tbl)
    tbl.columns[0].width = Inches(c0w)
    tbl.columns[1].width = Inches(c1w)
    c0, c1 = tbl.rows[0].cells
    c0.width = Inches(c0w)
    c1.width = Inches(c1w)

    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p0, line=1.0, after=0)
    _set_font(p0.add_run(ids), size=BODY_PT)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _spacing(p1, line=1.0, after=0)
    _set_font(p1.add_run(title), size=BODY_PT)

    tail = header.add_paragraph()
    _spacing(tail, line=1.0, after=0, before=0)
    _set_font(tail.add_run(""), size=1)


def _clear_hf(container):
    for p in list(container.paragraphs):
        p._element.getparent().remove(p._element)
    for t in list(container.tables):
        t._element.getparent().remove(t._element)
    return container.add_paragraph()


def _build_footer(section):
    section.footer.is_linked_to_previous = False
    p = _clear_hf(section.footer)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _spacing(p, line=1.0)
    _add_field(p, "PAGE", placeholder="1")


def _blank_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear_hf(section.header)
    _clear_hf(section.footer)


# --------------------------------------------------------------------------- #
# inline markdown
# --------------------------------------------------------------------------- #
_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def add_inline(par, text, size=BODY_PT, base_bold=False):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            _set_font(r, size=size, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            _set_font(r, size=size, bold=base_bold, italic=True)
        else:
            r = par.add_run(part)
            _set_font(r, size=size, bold=base_bold)


# --------------------------------------------------------------------------- #
# body document construction (front matter + chapters)
# --------------------------------------------------------------------------- #
def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.line_spacing = LINE
    for name, size in (
        ("Normal", BODY_PT),
        ("Heading 1", H1_PT),
        ("Heading 2", H2_PT),
        ("Heading 3", H3_PT),
    ):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = name.startswith("Heading")
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        st.paragraph_format.line_spacing = LINE
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def setup_geometry(doc):
    s = doc.sections[0]
    s.page_width = Twips(11907)
    s.page_height = Twips(16839)
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1.25)
    s.right_margin = Inches(1)
    s.header_distance = Inches(0.5)
    s.footer_distance = Inches(0.5)


def enable_update_fields(doc):
    try:
        settings = doc.settings.element
    except Exception:
        return
    for el in settings.findall(qn("w:updateFields")):
        settings.remove(el)
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def fm_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, line=LINE, after=12)
    _set_font(p.add_run(text), size=18, bold=True)


def fm_sub_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, line=LINE, before=6, after=12)
    _set_font(p.add_run(text), size=18, bold=True)


def body_para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=BODY_PT, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    _spacing(p, line=LINE)
    add_inline(p, text, size=size)
    if italic:
        for r in p.runs:
            r.font.italic = True
    return p


def render_markdown_table(doc, rows, caption=None, table_no=None):
    if caption and table_no:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(cap, line=1.0, before=6, after=2)
        _add_tc_field(cap, f"Table {table_no}: {caption}", "T")
        r = cap.add_run(f"Table {table_no}: {caption}")
        _set_font(r, size=BODY_PT, bold=True)

    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    tbl.autofit = True
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = tbl.cell(ri, ci)
            txt = row[ci] if ci < len(row) else ""
            par = cell.paragraphs[0]
            _spacing(par, line=1.0, after=0)
            add_inline(par, txt, size=BODY_PT, base_bold=(ri == 0))
    sp = doc.add_paragraph()
    _spacing(sp, line=0.5)


def render_figure(doc, no, caption, description):
    ph = doc.add_paragraph()
    ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(ph, line=1.0, before=6, after=2)
    _set_font(
        ph.add_run(f"(Insert Figure {no} here)"),
        size=BODY_PT,
        italic=True,
        color=(128, 128, 128),
    )
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(cap, line=1.0, after=2)
    _add_tc_field(cap, f"Figure {no}: {caption}", "F")
    _set_font(cap.add_run(f"Figure {no}: {caption}"), size=BODY_PT, bold=True)
    if description:
        desc = doc.add_paragraph()
        desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _spacing(desc, line=LINE, after=8)
        _set_font(desc.add_run(description), size=BODY_PT, italic=True)


_TBL_ROW = re.compile(r"^\s*\|.*\|\s*$")


def render_body(doc, text, meta):
    lines = text.splitlines()
    i = 0
    pending_caption = None
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        i += 1

        if not line:
            continue

        m = re.match(r"^\[\[FIG(?:URE)? ([^\|]+)\|(.+?)\]\]$", line)
        if m:
            no = m.group(1).strip()
            caption = m.group(2).strip()
            desc_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                desc_lines.append(lines[i].strip()[1:].strip())
                i += 1
            render_figure(doc, no, caption, " ".join(desc_lines))
            continue

        m = re.match(r"^\[\[TABLE ([^\|]+)\|(.+?)\]\]$", line)
        if m:
            pending_caption = (m.group(1).strip(), m.group(2).strip())
            continue

        if re.match(r"^\[\[PAGEBREAK\]\]$", line):
            _page_break(doc)
            continue

        if re.match(r"^\[\[FM:PAPER_PUBLICATION\]\]$", line):
            fm_sub_title(doc, "Paper Publication")
            continue

        m = re.match(r"^# (.+)$", line)
        if m and not line.startswith("##"):
            doc.add_heading(m.group(1), level=1)
            continue
        m = re.match(r"^## (.+)$", line)
        if m and not line.startswith("###"):
            doc.add_heading(m.group(1), level=2)
            continue
        m = re.match(r"^### (.+)$", line)
        if m:
            doc.add_heading(m.group(1), level=3)
            continue

        if _TBL_ROW.match(raw):
            rows = []
            while i <= len(lines) and _TBL_ROW.match(lines[i - 1]):
                cells = [c.strip() for c in lines[i - 1].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            if pending_caption:
                render_markdown_table(
                    doc, rows, caption=pending_caption[1], table_no=pending_caption[0]
                )
            else:
                render_markdown_table(doc, rows)
            pending_caption = None
            continue

        if re.match(r"^[-*] (.+)$", line):
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _spacing(p, line=LINE)
            add_inline(p, line[2:])
            continue

        m = re.match(r"^\[(\d+)\] (.+)$", line)
        if m:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _spacing(p, line=LINE)
            pf = p.paragraph_format
            pf.left_indent = Inches(0.5)
            pf.first_line_indent = Inches(-0.5)
            add_inline(p, f"[{m.group(1)}] {m.group(2)}")
            continue

        if line.startswith(">"):
            body_para(doc, line[1:].strip(), italic=True)
            continue

        body_para(doc, line)


def _set_cell_margins(table, top=0, bottom=0, left=40, right=40):
    tblpr = table._tbl.tblPr
    for el in tblpr.findall(qn("w:tblCellMar")):
        tblpr.remove(el)
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def render_abbreviations(doc, text):
    rows = []
    for line in text.splitlines():
        s = line.strip()
        if not _TBL_ROW.match(s):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.style = "Table Grid"
    tbl.autofit = True
    _set_cell_margins(tbl)
    for ri, row in enumerate(rows):
        for ci in range(2):
            par = tbl.cell(ri, ci).paragraphs[0]
            _spacing(par, line=1.0, before=0, after=0)
            add_inline(par, row[ci] if ci < len(row) else "", size=11, base_bold=(ri == 0))


def build_body_document(meta, fm, body_md):
    doc = Document()
    configure_styles(doc)
    setup_geometry(doc)

    if fm.get("acknowledgement"):
        fm_title(doc, "ACKNOWLEDGEMENT")
        for para in fm["acknowledgement"].split("\n\n"):
            for line in para.splitlines():
                if line.strip():
                    body_para(doc, line.strip())
        _page_break(doc)

    if fm.get("abstract"):
        fm_title(doc, "ABSTRACT")
        for para in fm["abstract"].split("\n\n"):
            for line in para.splitlines():
                if line.strip():
                    body_para(doc, line.strip())
        _page_break(doc)

    fm_title(doc, "LIST OF FIGURES")
    p = doc.add_paragraph()
    _spacing(p, line=LINE)
    _add_field(p, r"TOC \f F \h")
    _page_break(doc)

    fm_title(doc, "LIST OF TABLES")
    p = doc.add_paragraph()
    _spacing(p, line=LINE)
    _add_field(p, r"TOC \f T \h")
    _page_break(doc)

    fm_title(doc, "LIST OF ABBREVIATIONS")
    if fm.get("abbreviations"):
        render_abbreviations(doc, fm["abbreviations"])
    _page_break(doc)

    fm_title(doc, "TABLE OF CONTENTS")
    p = doc.add_paragraph()
    _spacing(p, line=LINE)
    _add_field(p, r'TOC \o "1-3" \h \z \u')

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    _pgnum(sec, fmt="decimal", start=1)

    render_body(doc, body_md, meta)
    enable_update_fields(doc)
    return doc


# --------------------------------------------------------------------------- #
# section repair (docxcompose merges sections; we split cover / certificate /
# front matter / chapters back into four sections)
# --------------------------------------------------------------------------- #
def _clean_sectpr(el):
    el = deepcopy(el)
    for tag in ("w:headerReference", "w:footerReference", "w:pgNumType"):
        for c in el.findall(qn(tag)):
            el.remove(c)
    return el


def _insert_section_break_before(anchor, sectpr):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), "20")
    sp.set(qn("w:lineRule"), "exact")
    ppr.append(sp)
    ppr.append(sectpr)
    p.append(ppr)
    anchor.addprevious(p)
    return p


def _fix_sections(doc, cover_sectpr, cert_sectpr, ids, title):
    body = doc.element.body
    cover_sectpr = _clean_sectpr(cover_sectpr)
    cert_sectpr = _clean_sectpr(cert_sectpr)

    first_tbl = None
    for child in body:
        if child.tag == qn("w:tbl"):
            first_tbl = child
            break
    if first_tbl is not None:
        _insert_section_break_before(first_tbl, cover_sectpr)

    ack = None
    for p in body.findall(qn("w:p")):
        if "".join(t.text or "" for t in p.iter(qn("w:t"))).strip() == "ACKNOWLEDGEMENT":
            ack = p
            break
    if ack is not None:
        _insert_section_break_before(ack, cert_sectpr)

    secs = doc.sections
    if len(secs) >= 2:
        _blank_header_footer(secs[0])
        _blank_header_footer(secs[1])
    for idx in range(2, len(secs)):
        _pgnum(secs[idx], fmt="lowerRoman" if idx == 2 else "decimal", start=1)
        _build_header(secs[idx], ids, title, _content_width_in(secs[idx]))
        _build_footer(secs[idx])
    return len(secs)


# --------------------------------------------------------------------------- #
# reading locked files on Windows
# --------------------------------------------------------------------------- #
def _read_bytes(path):
    try:
        with open(path, "rb") as fh:
            return fh.read()
    except PermissionError:
        pass
    if os.name != "nt":
        raise
    GENERIC_READ = 0x80000000
    OPEN_EXISTING = 3
    FILE_SHARE = 0x01 | 0x02 | 0x04
    INVALID = ctypes.c_void_p(-1).value
    k = ctypes.windll.kernel32
    handle = k.CreateFileW(
        ctypes.c_wchar_p(path),
        GENERIC_READ,
        FILE_SHARE,
        None,
        OPEN_EXISTING,
        0,
        None,
    )
    if handle in (INVALID, 0xFFFFFFFFFFFFFFFF):
        raise OSError(f"Cannot open file (locked?): {path}")
    try:
        chunks = []
        buf = ctypes.create_string_buffer(1 << 20)
        read = wintypes.DWORD()
        while True:
            ok = k.ReadFile(handle, buf, len(buf), ctypes.byref(read), None)
            if not ok or read.value == 0:
                break
            chunks.append(buf.raw[: read.value])
        return b"".join(chunks)
    finally:
        k.CloseHandle(handle)


# --------------------------------------------------------------------------- #
# forcing Times New Roman everywhere
# --------------------------------------------------------------------------- #
_RFONTS_TAG = re.compile(r"<w:rFonts\b[^>]*?/?>")


def _fix_rfonts_tag(tag):
    theme_attrs = ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme")
    explicit = ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia")
    for a in theme_attrs:
        tag = re.sub(rf'\s*{a}="[^"]*"', "", tag)
    if "Symbol" in tag or "Wingdings" in tag:
        return tag
    for a in explicit:
        tag = re.sub(rf'\s*{a}="[^"]*"', "", tag)
    closing = "/>" if tag.rstrip().endswith("/>") else ">"
    inner = tag.rstrip()[: -len(closing)]
    inner += (
        f' w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}" w:eastAsia="{FONT}"'
    )
    return inner + closing


def _fix_styles_xml(xml):
    xml = _RFONTS_TAG.sub(lambda m: _fix_rfonts_tag(m.group(0)), xml)

    def fix_default(m):
        block = m.group(0)
        if "<w:rPrDefault>" in block and "w:rFonts" not in block:
            rpr = re.search(r"<w:rPr>", block)
            if rpr:
                block = (
                    block[: rpr.end()]
                    + f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}" '
                    f'w:eastAsia="{FONT}"/>'
                    + block[rpr.end():]
                )
        return block

    xml = re.sub(r"<w:docDefaults>.*?</w:docDefaults>", fix_default, xml, flags=re.S)
    return xml


def _fix_theme_xml(xml):
    for tag in ("majorFont", "minorFont"):
        xml = re.sub(
            rf'(<a:{tag}>.*?<a:latin typeface=")[^"]*(")',
            rf"\1{FONT}\2",
            xml,
            count=1,
            flags=re.S,
        )
    return xml


def _rewrite_xml_parts(path, transforms):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    for name, fn in transforms.items():
        if name in data:
            text = data[name].decode("utf-8")
            data[name] = fn(text).encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    os.replace(tmp, path)


def force_times_new_roman(path):
    _rewrite_xml_parts(
        path,
        {
            "word/styles.xml": _fix_styles_xml,
            "word/theme/theme1.xml": _fix_theme_xml,
        },
    )


# --------------------------------------------------------------------------- #
# parsing
# --------------------------------------------------------------------------- #
def split_front_matter(text):
    m = re.search(r"(?:^|\n)---[ \t]*\n(.*?)\n---[ \t]*(?:\n|$)", text, re.S)
    if not m:
        return {}, text
    meta_text = m.group(1)
    body = text[m.end():]
    if yaml is None:
        raise SystemExit("PyYAML is required. pip install pyyaml")
    return yaml.safe_load(meta_text) or {}, body


def extract_fm(body):
    fm = {}
    for name in ("ACKNOWLEDGEMENT", "ABSTRACT", "ABBREVIATIONS"):
        pat = re.compile(rf"\[\[FM:{name}\]\](.*?)(?=\[\[FM:|\[\[BODY\]\])", re.S)
        m = pat.search(body)
        if m:
            fm[name.lower()] = m.group(1).strip()
    m = re.search(r"\[\[BODY\]\](.*)$", body, re.S)
    body_md = m.group(1).strip() if m else body
    return fm, body_md


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def generate(content_path, out_path, cover_path, cert_path, update_fields=True):
    with open(content_path, encoding="utf-8") as fh:
        raw = fh.read()
    meta, rest = split_front_matter(raw)
    fm, body_md = extract_fm(rest)

    ids = meta.get("student_id_header") or ", ".join(s["id"] for s in meta.get("students", []))
    title = meta.get("title", "")

    body_doc = build_body_document(meta, fm, body_md)

    cover_bytes = _read_bytes(cover_path)
    cert_bytes = _read_bytes(cert_path)

    cover_doc = Document(io.BytesIO(cover_bytes))
    cert_doc = Document(io.BytesIO(cert_bytes))
    cover_sectpr = deepcopy(cover_doc.element.body.find(qn("w:sectPr")))
    cert_sectpr = deepcopy(cert_doc.element.body.find(qn("w:sectPr")))

    composer = Composer(cover_doc)
    composer.append(cert_doc)
    composer.append(body_doc)
    merged = io.BytesIO()
    composer.save(merged)
    merged.seek(0)

    doc = Document(merged)
    _fix_sections(doc, cover_sectpr, cert_sectpr, ids, title)

    if not update_fields:
        for el in doc.settings.element.findall(qn("w:updateFields")):
            doc.settings.element.remove(el)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    force_times_new_roman(out_path)
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser(description="Generate the 4IT31 Project I report .docx")
    ap.add_argument("content", help="path to content.md")
    ap.add_argument("-o", "--output", required=True, help="output .docx path")
    ap.add_argument("--cover-template", default=DEFAULT_COVER)
    ap.add_argument("--certificate-template", default=DEFAULT_CERT)
    ap.add_argument("--no-toc-update", action="store_true")
    args = ap.parse_args(argv)

    out = generate(
        args.content,
        args.output,
        args.cover_template,
        args.certificate_template,
        update_fields=not args.no_toc_update,
    )
    print(f"Generated: {out}")


if __name__ == "__main__":
    sys.exit(main())
